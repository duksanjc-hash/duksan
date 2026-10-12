# -*- coding: utf-8 -*-
"""생산 RAW 일정/공장 반영 도구 (새 양식 K/M열 대응).
필요 패키지: openpyxl.  파일의 셀 스타일과 기타 시트를 유지하는 방향으로 저장합니다.
"""
from __future__ import annotations
import csv
from collections import defaultdict
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

HEADER = {'B': '도면번호', 'K': '제작예정일', 'M': '제작공장'}


def load_requests(path):
    items = defaultdict(int)
    for encoding in ('utf-8-sig', 'cp949', 'utf-8'):
        try:
            with open(path, 'r', newline='', encoding=encoding) as f:
                for n, row in enumerate(csv.reader(f), 1):
                    if not row or all(not x.strip() for x in row):
                        continue
                    if len(row) < 2:
                        raise ValueError(f'CSV {n}행: 부재번호와 수량 2개 열이 필요합니다.')
                    code = row[0].strip()
                    if n == 1 and code in ('부재번호', '도면번호'):
                        continue
                    try:
                        qty = int(row[1].strip())
                    except ValueError:
                        raise ValueError(f'CSV {n}행: 수량이 정수가 아닙니다: {row[1]}')
                    if qty < 0:
                        raise ValueError(f'CSV {n}행: 음수 수량은 처리하지 않습니다.')
                    if not code:
                        raise ValueError(f'CSV {n}행: 부재번호가 비어 있습니다.')
                    items[code] += qty
            return dict(items)
        except UnicodeDecodeError:
            items.clear()
            continue
    raise ValueError('CSV 파일의 인코딩을 확인할 수 없습니다.')


def txt(value):
    return '' if value is None else str(value).strip()


def process(raw_file, csv_file, dest_file, production_date, plant,
            overwrite_plant=False, placeholder='17', apply=False):
    """K열과 M열이 모두 빈 행에만 적용; 하나라도 입력되어 있으면 행 전체 보호.

    CSV 수량은 해당 부재번호에 신규 배정할 개수로 해석합니다.
    변경 대상 검증을 위해 첫 실행은 미리보기를 권장합니다.
    """
    requested = load_requests(csv_file)
    wb = load_workbook(raw_file)
    if '제작예정일' not in wb.sheetnames:
        raise ValueError('RAW 파일에 제작예정일 시트가 없습니다.')
    sheet = wb['제작예정일']
    for col, header in HEADER.items():
        if txt(sheet[f'{col}1'].value) != header:
            raise ValueError(f'{col}1 제목이 "{header}"가 아닙니다. 변경된 RAW 양식을 확인해주세요.')
    groups = defaultdict(list)
    for row in range(2, sheet.max_row + 1):
        code = txt(sheet.cell(row, 2).value)
        if code:
            groups[code].append(row)

    report = []
    edits = []
    for code, qty in requested.items():
        rows = groups.get(code, [])
        already = sum(bool(txt(sheet.cell(r, 11).value)) for r in rows)
        protected_plant = sum(bool(txt(sheet.cell(r, 13).value)) for r in rows)
        # 기존값이 어느 한 열에라도 있으면 K/M열 모두 수정 금지.
        eligible = [r for r in rows if not txt(sheet.cell(r, 11).value)
                    and not txt(sheet.cell(r, 13).value)]
        selected = eligible[:qty]
        for r in selected:
            edits.append((r, production_date, plant))
        changed = len(selected)
        if qty == 0:
            status = '수량 0 (제외)'
        elif not rows:
            status = 'RAW 부재번호 없음'
        elif len(eligible) < qty:
            status = '입력 가능 행 부족 (기존값 보호)'
        else:
            status = '반영 가능'
        report.append((code, qty, len(rows), already, len(eligible), changed, protected_plant, status))

    if apply:
        for r, date, p in edits:
            sheet.cell(r, 11).value = date
            sheet.cell(r, 11).number_format = 'yyyy-mm-dd'
            if p is not None:
                sheet.cell(r, 13).value = p

        report_name = '수량검증_더블체크'
        if report_name in wb.sheetnames:
            del wb[report_name]
        audit = wb.create_sheet(report_name)
        audit.append(['도면번호', 'CSV 요청수량', 'RAW 전체행', '기존 제작예정일 입력행',
                      '신규 입력 가능행', '이번 적용수량', '기존 제작공장 입력행', '검증결과'])
        for line in report:
            audit.append(line)
        for cell in audit[1]:
            cell.fill = PatternFill('solid', fgColor='17365D')
            cell.font = Font(color='FFFFFF', bold=True)
        for col, width in {'A': 22, 'B': 17, 'C': 14, 'D': 24,
                           'E': 20, 'F': 18, 'G': 25, 'H': 26}.items():
            audit.column_dimensions[col].width = width
        audit.freeze_panes = 'A2'
        audit.auto_filter.ref = audit.dimensions
        wb.save(dest_file)
    return report


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title('생산 RAW 자동 변경 | K열 제작예정일 / M열 제작공장')
        self.geometry('830x650')
        self.raw = tk.StringVar()
        self.csv = tk.StringVar()
        self.dest = tk.StringVar()
        self.date = tk.StringVar()
        self.plant = tk.StringVar(value='진천')
        self.create_form()

    def create_form(self):
        p = ttk.Frame(self, padding=18)
        p.pack(fill='both', expand=True)
        ttk.Label(p, text='생산 RAW 수정 프로그램 (변경 양식 전용)', font=('', 15, 'bold')).pack(anchor='w', pady=(0, 10))
        ttk.Label(p, text='기존 J/L열은 수정하지 않습니다. K열과 M열이 모두 빈 행만 대상입니다.').pack(anchor='w', pady=(0, 12))
        for label, var, filetype in (
            ('RAW Excel', self.raw, 'xlsx'), ('요청 CSV', self.csv, 'csv'), ('결과 저장', self.dest, 'save')):
            row = ttk.Frame(p); row.pack(fill='x', pady=4)
            ttk.Label(row, text=label, width=12).pack(side='left')
            ttk.Entry(row, textvariable=var).pack(side='left', fill='x', expand=True, padx=5)
            ttk.Button(row, text='찾아보기', command=lambda v=var, t=filetype: self.browse(v, t)).pack(side='left')
        row = ttk.Frame(p); row.pack(fill='x', pady=12)
        ttk.Label(row, text='제작예정일', width=12).pack(side='left')
        ttk.Entry(row, textvariable=self.date, width=18).pack(side='left', padx=5)
        ttk.Label(row, text='예: 2026-10-12').pack(side='left', padx=(2,20))
        ttk.Label(row, text='제작공장').pack(side='left')
        ttk.Entry(row, textvariable=self.plant, width=16).pack(side='left', padx=5)
        actions = ttk.Frame(p); actions.pack(fill='x', pady=5)
        ttk.Button(actions, text='1. 사전 검증 (저장 안 함)', command=lambda: self.run(False)).pack(side='left', padx=(0,8))
        ttk.Button(actions, text='2. 결과 파일 생성', command=lambda: self.run(True)).pack(side='left')
        self.log = tk.Text(p, height=19, wrap='none')
        self.log.pack(fill='both', expand=True, pady=(10,0))
        self.log.insert('end', 'CSV 수량 = 이번에 새로 제작예정일을 입력할 수량으로 처리합니다.\n'
                               'K열 또는 M열에 값이 하나라도 있으면 해당 행은 두 열 모두 수정하지 않습니다.\n')

    def browse(self, var, kind):
        if kind == 'save':
            val = filedialog.asksaveasfilename(defaultextension='.xlsx',
                  filetypes=[('Excel', '*.xlsx')], initialfile='생산RAW_수정결과.xlsx')
        else:
            val = filedialog.askopenfilename(filetypes=[('Excel', '*.xlsx')] if kind == 'xlsx' else [('CSV','*.csv')])
        if val:
            var.set(val)
            if kind == 'xlsx' and not self.dest.get():
                self.dest.set(str(Path(val).with_name(Path(val).stem+'_수정결과.xlsx')))

    def run(self, apply):
        try:
            if not self.raw.get() or not self.csv.get():
                raise ValueError('RAW Excel과 CSV 파일을 선택해주세요.')
            if not self.plant.get().strip():
                raise ValueError('제작공장을 입력해주세요.')
            try:
                d = datetime.strptime(self.date.get().strip(), '%Y-%m-%d')
            except ValueError:
                raise ValueError('제작예정일은 YYYY-MM-DD 형식으로 입력해주세요.')
            if apply and not self.dest.get():
                raise ValueError('결과 저장 경로를 지정해주세요.')
            if apply and Path(self.dest.get()).resolve() == Path(self.raw.get()).resolve():
                raise ValueError('원본 보호를 위해 결과 파일은 다른 이름으로 저장해주세요.')
            report = process(self.raw.get(), self.csv.get(), self.dest.get(), d,
                             self.plant.get().strip(), apply=apply)
            totals = [sum(line[i] for line in report) for i in (1,5,6)]
            issue = [x for x in report if x[7] not in ('반영 가능', '수량 0 (제외)')]
            self.log.delete('1.0', 'end')
            self.log.insert('end', f'CSV 요청 합계: {totals[0]}건 / 신규 예정일 적용: {totals[1]}건 / 기존 공장값 보존행: {totals[2]}건\n')
            self.log.insert('end', f'검증 안내: {len(issue)}종 (RAW 없음, 입력 가능 행 부족)\n\n')
            for x in issue[:120]:
                self.log.insert('end', f'{x[0]} | 요청 {x[1]} | 가능 {x[4]} | 적용 {x[5]} | {x[7]}\n')
            if len(issue)>120:
                self.log.insert('end', f'... 이외 {len(issue)-120}건은 결과 검증 시트 확인\n')
            if apply:
                messagebox.showinfo('완료', f'결과 저장 완료\n{self.dest.get()}\n\n반드시 수량검증_더블체크 시트를 확인해주세요.')
        except Exception as exc:
            messagebox.showerror('처리 오류', str(exc))


if __name__ == '__main__':
    App().mainloop()
